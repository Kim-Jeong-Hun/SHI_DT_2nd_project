"""Read-only checks tying the variable dictionary to executable project rules.

The workbook still owns human-written descriptions. This module deliberately
does not decide which cells may be overwritten or modify the workbook.
"""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from .legacy import PROJECT_ROOT, legacy_module
from .modeling import ROLLING_BASELINE


DICTIONARY_PATH = Path("data/catalog/선박수주예측_변수_데이터사전.xlsx")
FEATURE_SHEET = "Feature_Dictionary"
FEATURE_HEADER = (
    "대분류", "중분류", "변수명(현행·후보)", "정의·단위", "벌크선",
    "컨테이너선", "가스 운반선", "유조선", "케미컬 탱커", "원천 ID",
    "시차·활용 상태", "제공기관·자료", "원천 URL 또는 확인 위치",
    "raw 파일", "processed 파일", "원천 정의·주의",
)


@dataclass(frozen=True)
class Issue:
    code: str
    location: str
    detail: str


def _rows(sheet, first_row: int, last_column: int):
    for number, cells in enumerate(
        sheet.iter_rows(min_row=first_row, max_col=last_column, values_only=True),
        start=first_row,
    ):
        if any(value is not None and value != "" for value in cells):
            yield number, cells


def check_source_metadata(rows) -> tuple[list[Issue], set[str]]:
    """Repeated source IDs are allowed, but their metadata must agree."""
    issues: list[Issue] = []
    first_by_id: dict[str, tuple[int, tuple]] = {}
    for number, row in rows:
        source_id = row[9]
        if not isinstance(source_id, str) or not source_id.strip():
            issues.append(Issue("source.id_missing", f"{FEATURE_SHEET}!J{number}", "원천 ID가 비어 있습니다."))
            continue
        metadata = tuple(row[11:16])
        for column_index, value in enumerate(metadata, start=11):
            if not isinstance(value, str) or not value.strip():
                column = _excel_column(column_index + 1)
                issues.append(Issue("source.metadata_missing", f"{FEATURE_SHEET}!{column}{number}",
                                    f"{source_id}: {FEATURE_HEADER[column_index]} 값이 비어 있습니다."))
        if source_id in first_by_id:
            first_row, first_metadata = first_by_id[source_id]
            if metadata != first_metadata:
                issues.append(Issue("source.metadata_conflict", f"{FEATURE_SHEET}!L{number}:P{number}",
                                    f"{source_id}: {first_row}행의 원천 정보와 다릅니다."))
        else:
            first_by_id[source_id] = (number, metadata)
    return issues, set(first_by_id)


def _check_header(sheet, expected: tuple[str, ...], issues: list[Issue]) -> None:
    actual = tuple(cell.value for cell in sheet[5][:len(expected)])
    if actual != expected:
        issues.append(Issue("dictionary.header_mismatch", f"{sheet.title}!A5:{_excel_column(len(expected))}5",
                            f"예상 헤더 {expected!r}, 실제 {actual!r}"))


def expected_flags(rules) -> dict[str, dict[str, str]]:
    """Derive availability from the same rules used to create master tables."""
    expected: dict[str, dict[str, str]] = {}
    audit_names = {"기준월", rules.TARGET_UNIT_COLUMN}
    for ship_type in rules.SHIP_TYPES:
        features = set(rules.feature_columns_for(ship_type))
        for name in audit_names:
            expected.setdefault(name, {})[ship_type] = "R"
        expected.setdefault(rules.TARGET_COLUMN, {})[ship_type] = "Y"
        expected.setdefault(ROLLING_BASELINE, {})[ship_type] = "B"
        for name in features:
            expected.setdefault(name, {})[ship_type] = (
                "R" if name in rules.PROCESSED_ORDER_LAG_COLUMNS.values() else "O"
            )
    return expected


def check_feature_rows(rows, ship_columns, expected) -> tuple[list[Issue], int]:
    issues: list[Issue] = []
    seen: set[str] = set()
    candidates = 0
    for number, row in rows:
        name = row[2]
        location = f"{FEATURE_SHEET}!C{number}"
        if not isinstance(name, str) or not name.strip():
            issues.append(Issue("feature.name_missing", location, "변수명이 비어 있습니다."))
            continue
        if name in seen:
            issues.append(Issue("feature.name_duplicate", location, f"중복 변수명: {name}"))
        seen.add(name)
        flags = {ship: row[column] for ship, column in ship_columns.items()}
        if name not in expected:
            # 폐기된 과거 변수는 감사 이력으로 남겨도 신규 도입 후보로 세지 않는다.
            if not str(row[10] or "").startswith("폐기"):
                candidates += 1
            for ship, flag in flags.items():
                if flag != "-":
                    issues.append(Issue("feature.unimplemented_flag", f"{FEATURE_SHEET}!{_excel_column(ship_columns[ship] + 1)}{number}",
                                        f"{name}: 코드에 없는 후보는 {ship}에서 '-'여야 합니다. 현재 {flag!r}"))
            continue
        for ship, flag in flags.items():
            wanted = expected[name].get(ship, "-")
            if flag != wanted:
                issues.append(Issue("feature.flag_mismatch", f"{FEATURE_SHEET}!{_excel_column(ship_columns[ship] + 1)}{number}",
                                    f"{name}/{ship}: 엑셀 {flag!r}, 코드 {wanted!r}"))
    for name in sorted(set(expected) - seen):
        issues.append(Issue("feature.current_missing", FEATURE_SHEET,
                            f"코드의 현행 변수가 사전에 없습니다: {name}"))
    return issues, candidates


def _excel_column(number: int) -> str:
    result = ""
    while number:
        number, digit = divmod(number - 1, 26)
        result = chr(65 + digit) + result
    return result


def _check_master_headers(root: Path, rules) -> list[Issue]:
    issues: list[Issue] = []
    for ship in rules.SHIP_TYPES:
        path = root / "data/master" / f"마스터테이블({ship}).csv"
        if not path.is_file():
            issues.append(Issue("master.file_missing", str(path), "마스터 CSV가 없습니다."))
            continue
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            actual = next(csv.reader(handle), [])
        wanted = [
            "기준월", rules.TARGET_COLUMN, rules.TARGET_UNIT_COLUMN,
            "수주정보시차(개월)", "수주정보가용근거",
            *rules.feature_columns_for(ship),
        ]
        if actual != wanted:
            missing = sorted(set(wanted) - set(actual))
            extra = sorted(set(actual) - set(wanted))
            detail = f"코드와 CSV 헤더 불일치; 누락={missing}, 추가={extra}"
            if not missing and not extra:
                detail += "; 컬럼 순서가 다릅니다."
            issues.append(Issue("master.schema_mismatch", str(path), detail))
    return issues


def _check_orders_raw_header(root: Path, issues: list[Issue]) -> int:
    """Keep the original orders schema in code, not in another Excel tab."""
    mapping = legacy_module("수주실적_전처리").COLUMN_MAP
    path = root / "data/raw/수주실적.csv"
    if not path.is_file():
        issues.append(Issue("orders_raw.file_missing", str(path), "수주실적 원자료가 없습니다."))
        return len(mapping)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        actual = next(csv.reader(handle), [])
    if actual != list(mapping):
        missing = sorted(set(mapping) - set(actual))
        extra = sorted(set(actual) - set(mapping))
        issues.append(Issue("orders_raw.schema_mismatch", str(path),
                            f"코드의 원본 컬럼 규칙과 불일치; 누락={missing}, 추가={extra}"))
    return len(mapping)


def _check_orders_processed_headers(root: Path, rules) -> list[Issue]:
    issues: list[Issue] = []
    for ship in rules.SHIP_TYPES:
        path = root / "data/processed" / ship / f"수주실적({ship}).csv"
        if not path.is_file():
            issues.append(Issue("orders_processed.file_missing", str(path), "수주실적 가공본이 없습니다."))
            continue
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            actual = next(csv.reader(handle), [])
        wanted = ["기준월", "선종", f"수주용량({rules.TARGET_UNITS[ship]})",
                  *rules.PROCESSED_ORDER_LAG_COLUMNS.values()]
        if actual != wanted:
            issues.append(Issue("orders_processed.schema_mismatch", str(path),
                                f"코드와 CSV 헤더 불일치; 예상={wanted}, 실제={actual}"))
    return issues


def audit_project(root: Path = PROJECT_ROOT, workbook_path: Path | None = None) -> dict:
    root = Path(root)
    path = Path(workbook_path) if workbook_path else root / DICTIONARY_PATH
    issues: list[Issue] = []
    if not path.is_file():
        return {"workbook": str(path), "summary": {}, "issues": [asdict(Issue(
            "dictionary.file_missing", str(path), "변수 사전 엑셀이 없습니다."))]}
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        required = {FEATURE_SHEET}
        for name in sorted(required - set(workbook.sheetnames)):
            issues.append(Issue("dictionary.sheet_missing", str(path), f"필수 탭 누락: {name}"))
        for name in sorted(set(workbook.sheetnames) - required):
            issues.append(Issue("dictionary.sheet_unexpected", str(path), f"활성 사전에 불필요한 탭: {name}"))
        if issues:
            return {"workbook": str(path), "summary": {}, "issues": [asdict(issue) for issue in issues]}

        rules = legacy_module("마스터테이블_생성")
        feature_sheet = workbook[FEATURE_SHEET]
        _check_header(feature_sheet, FEATURE_HEADER, issues)
        ship_columns = {ship: FEATURE_HEADER.index(ship) for ship in rules.SHIP_TYPES}
        feature_rows = list(_rows(feature_sheet, 6, len(FEATURE_HEADER)))
        source_issues, source_ids = check_source_metadata(feature_rows)
        issues.extend(source_issues)
        feature_issues, candidate_count = check_feature_rows(
            feature_rows, ship_columns, expected_flags(rules)
        )
        issues.extend(feature_issues)
        issues.extend(_check_master_headers(root, rules))
        issues.extend(_check_orders_processed_headers(root, rules))
        mapping_count = _check_orders_raw_header(root, issues)
        summary = {
            "feature_rows": len(feature_rows),
            "current_features": len(expected_flags(rules)),
            "unintroduced_candidates": candidate_count,
            "master_tables": len(rules.SHIP_TYPES),
            "code_order_mappings": mapping_count,
            "source_ids": len(source_ids),
        }
    finally:
        workbook.close()
    return {"workbook": str(path), "summary": summary,
            "issues": [asdict(issue) for issue in issues]}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="변수 사전과 현행 코드·마스터 CSV를 읽기 전용으로 대조합니다.")
    parser.add_argument("--workbook", type=Path, help="검사할 엑셀 경로(기본: data/catalog의 현행 파일)")
    parser.add_argument("--json", action="store_true", help="결과를 JSON으로 출력")
    args = parser.parse_args(argv)
    report = audit_project(workbook_path=args.workbook)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"변수 사전 점검: {report['workbook']}")
        labels = {
            "feature_rows": "사전 변수",
            "current_features": "코드상 현행 변수",
            "unintroduced_candidates": "미도입 후보",
            "master_tables": "마스터 CSV",
            "code_order_mappings": "코드상 원본 컬럼 매핑",
            "source_ids": "원천 ID",
        }
        for key, value in report["summary"].items():
            print(f"  {labels.get(key, key)}: {value}")
        if report["issues"]:
            for issue in report["issues"]:
                print(f"오류 [{issue['code']}] {issue['location']}: {issue['detail']}")
        else:
            print("일치: 검사 범위에서 코드·마스터·엑셀의 불일치가 없습니다.")
    return int(bool(report["issues"]))


if __name__ == "__main__":
    raise SystemExit(main())
