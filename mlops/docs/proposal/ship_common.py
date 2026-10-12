# 선박 수주 예측 서비스 — 프로젝트 기획서 (PPTX) · 2026-10-06 주제 확정
#  톤앤매너: ShipHub 결과보고서 최종본과 동일 — 삼성 블루(#024AD8/#296EF9) · 먹색 #111214 · 회색 #4C525A · 선 #DDE0E5
#            Pretendard(SIL OFL, 유료 폰트 금지 조건 충족) · 표제 Bold / 부제 Medium / 본문 Regular · 어절 단위 줄바꿈
#            공공 템플릿 시안: 상단 가는 선 + PROJECT PROPOSAL / DATE 머리말, 여백 넉넉, 장식 최소
#  수치 출처: Drive "Ship orders forcasting" — reports/2026_독립검증.json, scripts/분석/수주예측(선종).ipynb, data/catalog 변수사전
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

BLUE = RGBColor(0x02, 0x4A, 0xD8); BLUEB = RGBColor(0x29, 0x6E, 0xF9); BLUEL = RGBColor(0xEE, 0xF3, 0xFD); BLUEM = RGBColor(0xD6, 0xE3, 0xFB)
INK = RGBColor(0x11, 0x12, 0x14); GRAY = RGBColor(0x4C, 0x52, 0x5A); GRAYL = RGBColor(0x8A, 0x90, 0x98); LINE = RGBColor(0xDD, 0xE0, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); OK = RGBColor(0x1D, 0x7F, 0x4E); WARN = RGBColor(0xB4, 0x5F, 0x06); BGS = RGBColor(0xF6, 0xF7, 0xF9)
F_REG, F_MED, F_SEMI = "Pretendard", "Pretendard Medium", "Pretendard SemiBold"

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BL = prs.slide_layouts[6]
TOTAL = 18
state = {"n": 0}


def run_fmt(r, size, font=F_REG, bold=False, color=INK):
    r.font.size = Pt(size); r.font.name = font; r.font.bold = bold; r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def tb(s, x, y, w, h, paras, size=13, font=F_REG, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25, after=4):
    """paras: str | list[str | list[(text, opts)]]  — opts: size/font/bold/color"""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = spacing; para.space_after = Pt(after)
        runs = p if isinstance(p, list) else [(p, {})]
        for text, o in runs:
            r = para.add_run(); r.text = text
            run_fmt(r, o.get("size", size), o.get("font", font), o.get("bold", bold), o.get("color", color))
    return box


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0, radius=None):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(lw)
    r.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        r.adjustments[0] = radius
    return r


def hline(s, x, y, w, color=LINE, wt=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(wt); return c


def arrow(s, x1, y1, x2, y2, color=GRAYL, wt=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(wt)
    ln = c.line._get_or_add_ln(); te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def table(s, x, y, w, rows, colw, size=11, rowh=0.36, head_fill=BLUE, first_bold=False, emph_rows=()):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr)).table
    tblPr = t._tbl.tblPr; tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
    st = tblPr.find(qn("a:tableStyleId"))
    if st is not None:
        st.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid → 테두리는 직접
    tot = sum(colw)
    for j, cw in enumerate(colw):
        t.columns[j].width = Inches(w * cw / tot)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.035)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True; tf.paragraphs[0].text = ""
            p = tf.paragraphs[0]; r = p.add_run(); r.text = str(val)
            head = i == 0
            run_fmt(r, size, F_SEMI if head or (first_bold and j == 0) or i in emph_rows else F_REG, False, WHITE if head else (BLUE if i in emph_rows else INK))
            p.line_spacing = 1.15
            c.fill.solid(); c.fill.fore_color.rgb = head_fill if head else (BLUEL if i in emph_rows else WHITE)
            tcPr = c._tc.get_or_add_tcPr()
            for k, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):   # 테두리는 tcPr 맨 앞(채우기보다 먼저)이어야 적용됨
                ln = etree.Element(qn(side), w=str(Emu(Pt(0.75)))); tcPr.insert(k, ln)
                if side in ("a:lnL", "a:lnR"):
                    etree.SubElement(ln, qn("a:noFill"))
                else:
                    sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val="DDE0E5")
    return t


def page(sec_no, sec_name, title, lead):
    s = prs.slides.add_slide(BL); state["n"] += 1
    hline(s, 0.6, 0.42, 12.13, BLUE, 1.0)
    tb(s, 0.6, 0.52, 4, 0.25, "PROJECT PROPOSAL", 9, F_MED, color=GRAY)
    tb(s, 8.73, 0.52, 4, 0.25, "선박 수주 예측 서비스 · 3조 선(船)견지명", 9, F_MED, color=GRAY, align=PP_ALIGN.RIGHT)
    tb(s, 0.6, 0.95, 1.0, 0.5, sec_no, 26, F_REG, True, BLUE)
    tb(s, 1.35, 1.0, 9, 0.45, [[(sec_name + "  ", {"size": 12, "font": F_MED, "color": BLUE}), (title, {"size": 24, "bold": True})]], anchor=MSO_ANCHOR.BOTTOM)
    tb(s, 1.35, 1.58, 11.2, 0.5, lead, 14, F_MED, color=GRAY)
    tb(s, 11.73, 7.0, 1.0, 0.25, f"{state['n']:02d} / {TOTAL:02d}", 9, F_MED, color=GRAYL, align=PP_ALIGN.RIGHT)
    hline(s, 0.6, 6.92, 12.13, LINE, 0.75)
    return s


def card(s, x, y, w, h, big, label, sub=None, color=BLUE):
    rect(s, x, y, w, h, BGS, None, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    tb(s, x + 0.22, y + 0.18, w - 0.44, 0.6, big, 26, F_REG, True, color)
    tb(s, x + 0.22, y + 0.82, w - 0.44, 0.35, label, 12.5, F_SEMI)
    if sub:
        tb(s, x + 0.22, y + 1.15, w - 0.44, h - 1.25, sub, 11, color=GRAY, spacing=1.2)


def chip(s, x, y, w, text, fill=BLUE, color=WHITE, size=11):
    r = rect(s, x, y, w, 0.34, fill, None, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tb(s, x, y + 0.045, w, 0.28, text, size, F_SEMI, color=color, align=PP_ALIGN.CENTER)

